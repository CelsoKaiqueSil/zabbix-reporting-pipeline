"""Create an auditable Excel report from synthetic Zabbix metric samples."""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill

REQUIRED_COLUMNS = {"client", "host", "item_name", "timestamp", "value"}


def classify_metric(item_name: str) -> str:
    """Classify an item while prioritizing explicit memory terms over processor."""
    name = item_name.casefold()
    if any(term in name for term in ("memory", "memória", "memoria", "ram")):
        return "Memória"
    if any(term in name for term in ("cpu", "processor", "load")):
        return "CPU"
    if any(term in name for term in ("traffic", "tráfego", "trafego", "bits", "bytes")):
        return "Tráfego"
    return "Outros"


def load_metrics(path: str | Path) -> pd.DataFrame:
    frame = pd.read_csv(path)
    missing = REQUIRED_COLUMNS.difference(frame.columns)
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(sorted(missing))}")

    frame = frame.copy()
    input_rows = len(frame)
    frame['source_row'] = frame.index + 2
    frame["timestamp"] = pd.to_datetime(frame["timestamp"], errors="coerce", utc=True)
    frame["value"] = pd.to_numeric(frame["value"], errors="coerce")
    reasons = pd.Series('', index=frame.index)
    for column in ['client', 'host', 'item_name']:
        reasons.loc[frame[column].isna() | frame[column].astype(str).str.strip().eq('')] += 'missing ' + column + '; '
    reasons.loc[frame['timestamp'].isna()] += 'invalid timestamp; '
    reasons.loc[frame['value'].isna() | frame['value'].abs().eq(float('inf'))] += 'invalid numeric value; '
    frame['unit'] = frame['unit'].fillna('unspecified') if 'unit' in frame else 'unspecified'
    rejected = frame.loc[reasons.ne('')].copy()
    rejected['rejection_reason'] = reasons.loc[reasons.ne('')]
    frame = frame.loc[reasons.eq('')].copy()
    frame["metric_group"] = frame["item_name"].map(classify_metric)
    frame = frame.sort_values(["client", "host", "item_name", "timestamp"]).reset_index(drop=True)
    frame.attrs['quality'] = {'input_rows': input_rows, 'accepted_rows': len(frame), 'rejected_rows': len(rejected)}
    frame.attrs['rejected'] = rejected
    return frame


def summarize_peaks(frame: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    keys = ["client", "host", "item_name", "metric_group", 'unit']
    for key, group in frame.groupby(keys, dropna=False, sort=True):
        maximum = group["value"].max()
        peak_rows = group.loc[group["value"].eq(maximum)].sort_values("timestamp")
        rows.append(
            {
                "client": key[0],
                "host": key[1],
                "item_name": key[2],
                "metric_group": key[3],
                'unit': key[4],
                "average": group["value"].mean(),
                "maximum": maximum,
                "first_peak_at": peak_rows.iloc[0]["timestamp"],
                "peak_count": int(len(peak_rows)),
                "sample_count": int(len(group)),
            }
        )
    return pd.DataFrame(rows, columns=[*keys, 'average', 'maximum', 'first_peak_at', 'peak_count', 'sample_count'])


def _summary_by_client(frame: pd.DataFrame) -> pd.DataFrame:
    return (
        frame.groupby(["client", "metric_group", 'item_name', 'unit'], as_index=False)
        .agg(samples=("value", "size"), average=("value", "mean"), maximum=("value", "max"))
        .sort_values(["client", "metric_group", 'item_name', 'unit'])
    )


def _excel_safe(frame: pd.DataFrame) -> pd.DataFrame:
    """Return a copy whose datetime columns are compatible with Excel."""
    result = frame.copy()
    for column in result.columns:
        if isinstance(result[column].dtype, pd.DatetimeTZDtype):
            result[column] = result[column].dt.tz_convert("UTC").dt.tz_localize(None)
    return result


def build_report(input_path: str | Path, output_path: str | Path) -> Path:
    frame = load_metrics(input_path)
    peaks = summarize_peaks(frame)
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    control = pd.DataFrame(
        [
            {"control": "generated_at_utc", "value": datetime.now(timezone.utc).isoformat()},
            {"control": "source", "value": "synthetic_csv"},
            {"control": "rows", "value": len(frame)},
            {"control": "clients", "value": frame["client"].nunique()},
            {"control": "read_only", "value": True},
            *[{'control': key, 'value': value} for key, value in frame.attrs['quality'].items()],
        ]
    )

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        _summary_by_client(frame).to_excel(writer, sheet_name="Resumo", index=False)
        _excel_safe(peaks).to_excel(writer, sheet_name="Picos", index=False)
        for metric in ("CPU", "Memória", "Tráfego"):
            _excel_safe(frame.loc[frame["metric_group"].eq(metric)]).to_excel(
                writer, sheet_name=metric, index=False
            )
        control.to_excel(writer, sheet_name="Controle", index=False)
        _excel_safe(frame.attrs['rejected']).to_excel(writer, sheet_name='Rejeitados', index=False)
        for sheet in writer.book:
            sheet.freeze_panes = 'A2'
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.fill = PatternFill('solid', fgColor='147E57')
                cell.font = Font(bold=True, color='FFFFFF')
            from openpyxl.utils import get_column_letter
            for index in range(1, sheet.max_column + 1):
                sheet.column_dimensions[get_column_letter(index)].width = 24
    return output


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", default="data/sample_metrics.csv")
    parser.add_argument("--output", default="output/zabbix_report.xlsx")
    args = parser.parse_args()
    path = build_report(args.input, args.output)
    print(path)


if __name__ == "__main__":
    main()
